# -*- coding: utf-8 -*-
"""
Consulta de Bonos y Puntualidad — Transportes LIPU Guadalajara
================================================================

Cada semana se sube UN solo Excel que trae todas las pestañas originales:
BONOS, RENDIMIENTO, DMAS, PUNTUALIDAD, CONSULTA y ACUMULADO.

Funciones principales:
- Reproduce la lógica de la pestaña "CONSULTA" (VLOOKUP/MATCH contra
  BONOS y RENDIMIENTO).
- Busca también en DMAS y PUNTUALIDAD por número de nómina.
- Si la nómina buscada NO aparece ni en BONOS ni en RENDIMIENTO, la
  búsqueda recurre automáticamente a ACUMULADO y muestra esos datos.
- También se puede buscar por nombre (aunque no sepas el número).
- Consulta múltiple: varias nóminas a la vez.
- Tendencia: cómo ha variado lo pagado semana a semana.
- Exportar cualquier consulta a Excel.
- Validación automática del archivo al subirlo.
- Respaldo con un clic de toda la información guardada.
- Eliminar, a futuro, cualquier semana completa ya cargada.

Uso: ejecutar por medio de "Ejecutar_ConsultaBonos.bat"
"""

import io
import json
import re
import shutil
import unicodedata
import zipfile
from datetime import datetime, date, time
from pathlib import Path

import openpyxl
import pandas as pd
import streamlit as st


# ---------------------------------------------------------------------------
# 🔐 AUTENTICACIÓN
# ---------------------------------------------------------------------------

def _password_hash(password: str, salt: bytes) -> str:
    """Genera un hash PBKDF2-SHA256 de la contraseña."""
    import hashlib
    return hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        310_000,
    ).hex()


def verificar_password(password: str, almacenado: str) -> bool:
    """Verifica una contraseña almacenada como SALT_HEX:HASH_HEX."""
    import hashlib
    import hmac
    try:
        salt_hex, hash_hex = almacenado.split(":", 1)
        salt = bytes.fromhex(salt_hex)
        esperado = bytes.fromhex(hash_hex)
        obtenido = hashlib.pbkdf2_hmac(
            "sha256",
            password.encode("utf-8"),
            salt,
            310_000,
        )
        return hmac.compare_digest(obtenido, esperado)
    except (ValueError, TypeError):
        return False


def mostrar_login() -> bool:
    """Muestra el login y devuelve True únicamente después de autenticar."""
    if st.session_state.get("autenticado", False):
        return True

    st.markdown(
        """
        <style>
        .login-box {
            max-width: 460px;
            margin: 80px auto 0 auto;
            padding: 30px 34px;
            border: 1px solid rgba(128,128,128,.25);
            border-radius: 14px;
            box-shadow: 0 4px 18px rgba(0,0,0,.08);
        }
        .login-title {
            text-align: center;
            font-size: 28px;
            font-weight: 700;
            margin-bottom: 6px;
        }
        .login-subtitle {
            text-align: center;
            color: #777;
            margin-bottom: 25px;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="login-box">'
        '<div class="login-title">🔐 Consulta de Bonos LIPU</div>'
        '<div class="login-subtitle">Inicia sesión para continuar</div>'
        '</div>',
        unsafe_allow_html=True,
    )

    # Los campos se muestran fuera del HTML para conservar los controles nativos
    # y el manejo seguro de contraseñas de Streamlit.
    with st.form("login_form"):
        usuario = st.text_input("Usuario", autocomplete="username")
        contrasena = st.text_input(
            "Contraseña",
            type="password",
            autocomplete="current-password",
        )
        entrar = st.form_submit_button("🔓 Iniciar sesión", type="primary", use_container_width=True)

        if entrar:
            try:
                usuarios = st.secrets["auth"]["users"]
            except Exception:
                st.error(
                    "No está configurada la autenticación. "
                    "Agrega las credenciales en Secrets de Streamlit."
                )
                st.stop()

            usuario_limpio = usuario.strip()
            if usuario_limpio in usuarios and verificar_password(
                contrasena, str(usuarios[usuario_limpio])
            ):
                st.session_state["autenticado"] = True
                st.session_state["usuario"] = usuario_limpio
                st.rerun()
            else:
                st.error("Usuario o contraseña incorrectos.")

    st.stop()


def mostrar_boton_cerrar_sesion():
    """Muestra usuario actual y botón para cerrar sesión."""
    usuario = st.session_state.get("usuario", "")
    with st.sidebar:
        st.markdown("---")
        st.markdown(f"👤 **Usuario:** {usuario}")
        if st.button("🚪 Cerrar sesión", use_container_width=True):
            st.session_state["autenticado"] = False
            st.session_state.pop("usuario", None)
            st.rerun()


mostrar_login()
mostrar_boton_cerrar_sesion()

# ---------------------------------------------------------------------------
# Configuración y rutas
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "datos"
HIST_DIR = BASE_DIR / "historico"
ACUM_DIR = DATA_DIR / "acumulado_por_semana"
ACTUAL_PATH = DATA_DIR / "actual.xlsx"
META_PATH = DATA_DIR / "meta.json"

DATA_DIR.mkdir(exist_ok=True)
HIST_DIR.mkdir(exist_ok=True)
ACUM_DIR.mkdir(exist_ok=True)

SHEET_BONOS = "BONOS"
SHEET_RENDIMIENTO = "RENDIMIENTO"
SHEET_DMAS = "DMAS"
SHEET_PUNTUALIDAD = "PUNTUALIDAD"
SHEET_ACUMULADO = "ACUMULADO"
SHEET_CONSULTA = "CONSULTA"

ACUM_COLUMNAS = ["AGRUPADOR", "CONCEPTO", "CLAVE", "NOMBRE_COMPLETO", "IMPORTE", "PERIODO"]

st.set_page_config(page_title="Consulta de Bonos LIPU", layout="wide")

# ---------------------------------------------------------------------------
# Utilidades generales
# ---------------------------------------------------------------------------


def norm(value) -> str:
    """Normaliza un encabezado/texto para comparar sin importar espacios/mayúsculas."""
    if value is None:
        return ""
    return str(value).strip().upper()


def normalizar_texto(s) -> str:
    """Quita acentos y normaliza mayúsculas, para comparar nombres de personas."""
    if s is None:
        return ""
    s = str(s).strip().upper()
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def clave_orden_periodo(periodo):
    """Extrae el número de un texto tipo 'S34' para poder ordenar semanas."""
    m = re.search(r"\d+", str(periodo))
    return int(m.group()) if m else 0


def to_display(value):
    """Convierte valores de celda a texto legible, con 2 decimales para números."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        try:
            if pd.isna(value):
                return ""
        except Exception:
            pass
        return f"{float(value):,.2f}"
    return value


def find_header_row(ws, target_norm: str, max_row: int = 15):
    """Busca la fila de encabezados que contiene target_norm en cualquier columna."""
    top = min(max_row, ws.max_row)
    for r in range(1, top + 1):
        for c in range(1, ws.max_column + 1):
            if norm(ws.cell(row=r, column=c).value) == target_norm:
                return r
    return None


def build_header_map(ws, header_row: int):
    """Regresa: hmap {encabezado_normalizado: columna} y lista [(columna, encabezado_original)]."""
    hmap = {}
    columns_ordered = []
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        if val is None or str(val).strip() == "":
            continue
        key = norm(val)
        if key not in hmap:  # primer encabezado repetido gana (igual que MATCH de Excel)
            hmap[key] = c
        columns_ordered.append((c, val))
    return hmap, columns_ordered


def build_nomina_index(ws, header_row: int, nomina_col: int, multi: bool):
    """Índice número de nómina -> fila (o lista de filas si multi=True)."""
    idx = {}
    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=nomina_col).value
        if raw is None or str(raw).strip() == "":
            continue
        try:
            key = int(raw)
        except (ValueError, TypeError):
            key = str(raw).strip()
        if multi:
            idx.setdefault(key, []).append(r)
        else:
            idx.setdefault(key, r)  # primer registro gana, igual que VLOOKUP
    return idx


def get_val(sheet_data, row, label_norm: str, default=""):
    if sheet_data is None or row is None:
        return default
    col = sheet_data["hmap"].get(label_norm)
    if col is None:
        return default
    v = sheet_data["ws"].cell(row=row, column=col).value
    return default if v is None else v


def parse_nomina_key(text: str):
    text = text.strip()
    try:
        return int(text)
    except ValueError:
        return text


def etiqueta_bonita(norm_label: str) -> str:
    return norm_label.replace("_", " ").title()


# ---------------------------------------------------------------------------
# ACUMULADO (nómina oficial) — almacén persistente, un CSV por semana
#
# La pestaña ACUMULADO ya trae su propia columna "periodo" (ej. "S34"), así
# que esa etiqueta se usa tal cual para separar y guardar cada semana en
# /datos/acumulado_por_semana/{periodo}.csv. Se extrae automáticamente cada
# vez que se sube el Excel semanal completo; subir de nuevo la misma semana
# simplemente reemplaza ese archivo (no se duplica).
# ---------------------------------------------------------------------------


def extraer_acumulado_por_periodo(path_xlsx) -> dict:
    """Si el Excel tiene pestaña ACUMULADO, regresa {periodo: DataFrame}."""
    wb = openpyxl.load_workbook(path_xlsx, data_only=True)
    if SHEET_ACUMULADO not in wb.sheetnames:
        return {}
    ws = wb[SHEET_ACUMULADO]
    hr = find_header_row(ws, "CLAVE")
    if hr is None:
        return {}
    hmap, _ = build_header_map(ws, hr)
    if not {"CLAVE", "IMPORTE", "PERIODO"}.issubset(hmap.keys()):
        return {}

    filas = []
    for r in range(hr + 1, ws.max_row + 1):
        clave = ws.cell(row=r, column=hmap["CLAVE"]).value
        if clave is None or str(clave).strip() == "":
            continue
        fila = {}
        for col_norm in ACUM_COLUMNAS:
            c = hmap.get(col_norm)
            fila[col_norm] = ws.cell(row=r, column=c).value if c else None
        filas.append(fila)

    if not filas:
        return {}

    df = pd.DataFrame(filas, columns=ACUM_COLUMNAS)
    df["CLAVE"] = pd.to_numeric(df["CLAVE"], errors="coerce")
    df = df.dropna(subset=["CLAVE"])
    df["CLAVE"] = df["CLAVE"].astype(int)

    grupos = {}
    for periodo, sub in df.groupby("PERIODO"):
        etiqueta = str(periodo).strip()
        if etiqueta:
            grupos[etiqueta] = sub.reset_index(drop=True)
    return grupos


def guardar_acumulado(grupos: dict):
    """Guarda (o reemplaza) el CSV de cada semana encontrada."""
    for periodo, df in grupos.items():
        destino = ACUM_DIR / f"{periodo}.csv"
        df.to_csv(destino, index=False, encoding="utf-8-sig")


def listar_periodos_acumulado():
    return sorted(ACUM_DIR.glob("*.csv"))


def buscar_acumulado_por_nomina(nomina) -> pd.DataFrame:
    """Junta, de todas las semanas guardadas, las filas de esa nómina."""
    if not isinstance(nomina, int):
        return pd.DataFrame()
    coincidencias = []
    for archivo in listar_periodos_acumulado():
        try:
            df = pd.read_csv(archivo, encoding="utf-8-sig")
        except Exception:
            continue
        df["CLAVE"] = pd.to_numeric(df["CLAVE"], errors="coerce")
        encontrado = df[df["CLAVE"] == nomina]
        if not encontrado.empty:
            coincidencias.append(encontrado)
    if not coincidencias:
        return pd.DataFrame()
    return pd.concat(coincidencias, ignore_index=True)


def mostrar_tabla_acumulado(df_acum: pd.DataFrame):
    """Renderiza el detalle de ACUMULADO agrupado por semana, con totales."""
    for periodo in sorted(df_acum["PERIODO"].astype(str).unique(), key=clave_orden_periodo):
        sub = df_acum[df_acum["PERIODO"].astype(str) == periodo][
            ["AGRUPADOR", "CONCEPTO", "IMPORTE"]
        ].copy()
        total_periodo = sub["IMPORTE"].sum()
        sub = sub.rename(columns={"AGRUPADOR": "Agrupador", "CONCEPTO": "Concepto", "IMPORTE": "Importe"})
        sub["Importe"] = sub["Importe"].apply(to_display)
        st.markdown(f"**Semana {periodo} — Total: ${total_periodo:,.2f}**")
        st.table(sub)
    if df_acum["PERIODO"].astype(str).nunique() > 1:
        st.markdown(f"**Total en todas las semanas mostradas: ${df_acum['IMPORTE'].sum():,.2f}**")


# ---------------------------------------------------------------------------
# Índice de nombres (para poder buscar por nombre, no solo por nómina)
# ---------------------------------------------------------------------------


def construir_indice_nombres(bonos, dmas) -> dict:
    """Regresa {nómina: nombre}, combinando BONOS, DMAS y ACUMULADO."""
    indice = {}
    if bonos:
        for nomina_val, fila in bonos["idx"].items():
            if isinstance(nomina_val, int):
                nombre = get_val(bonos, fila, "NOMBRE")
                if nombre:
                    indice.setdefault(nomina_val, str(nombre).strip())
    if dmas:
        for nomina_val, fila in dmas["idx"].items():
            if isinstance(nomina_val, int) and nomina_val not in indice:
                nombre = get_val(dmas, fila, "DRIVER_NAME")
                if nombre:
                    indice.setdefault(nomina_val, str(nombre).strip())
    for archivo in listar_periodos_acumulado():
        try:
            df = pd.read_csv(archivo, encoding="utf-8-sig", usecols=["CLAVE", "NOMBRE_COMPLETO"])
        except Exception:
            continue
        for _, fila in df.drop_duplicates(subset=["CLAVE"]).iterrows():
            try:
                clave = int(fila["CLAVE"])
            except (ValueError, TypeError):
                continue
            if clave not in indice and pd.notna(fila["NOMBRE_COMPLETO"]):
                indice.setdefault(clave, str(fila["NOMBRE_COMPLETO"]).strip())
    return indice


def buscar_por_nombre(indice_nombres: dict, texto: str):
    """Regresa lista de (nómina, nombre) cuyo nombre contiene 'texto' (sin acentos/mayúsculas)."""
    texto_norm = normalizar_texto(texto)
    coincidencias = [
        (nom, nombre) for nom, nombre in indice_nombres.items() if texto_norm in normalizar_texto(nombre)
    ]
    coincidencias.sort(key=lambda x: x[1])
    return coincidencias


# ---------------------------------------------------------------------------
# Validación del archivo al subirlo
# ---------------------------------------------------------------------------


def validar_archivo(path_xlsx):
    """Revisa que estén las pestañas y columnas clave que la app necesita."""
    resultados = []
    wb = openpyxl.load_workbook(path_xlsx, data_only=True)

    hojas_esperadas = [SHEET_BONOS, SHEET_RENDIMIENTO, SHEET_DMAS, SHEET_PUNTUALIDAD, SHEET_CONSULTA, SHEET_ACUMULADO]
    for hoja in hojas_esperadas:
        if hoja not in wb.sheetnames:
            resultados.append({"seccion": hoja, "estado": "falta", "detalle": "La pestaña no existe en el archivo."})
        else:
            resultados.append({"seccion": hoja, "estado": "ok", "detalle": "Pestaña encontrada."})

    columnas_esperadas = {
        SHEET_BONOS: [
            "CLAVE", "NOMBRE", "SINIESTROS", "QUEJAS", "SCORE OBTENIDO", "BONO SCORE", "PUNTUALIDAD",
            "BONO LOGUEO", "BONO SINIESTROS", "BONO QUEJAS", "BONO SCORE1", "BONO LOGUEO1", "TOTAL B.DESEMPEÑO",
        ],
        SHEET_RENDIMIENTO: ["CLAVE", "META", "REDIMIENTO OBTENIDO", "TOTAL A PAGAR"],
        SHEET_DMAS: ["NOMINA"],
        SHEET_PUNTUALIDAD: ["NO. DE NOMINA", "START_DATE", "PUNTUALIDAD", "ESTATUS", "SEMANA", "TURNO"],
        SHEET_ACUMULADO: ACUM_COLUMNAS,
    }
    claves_por_hoja = {
        SHEET_BONOS: "CLAVE",
        SHEET_RENDIMIENTO: "CLAVE",
        SHEET_DMAS: "NOMINA",
        SHEET_PUNTUALIDAD: "NO. DE NOMINA",
        SHEET_ACUMULADO: "CLAVE",
    }

    for hoja, columnas in columnas_esperadas.items():
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        hr = find_header_row(ws, claves_por_hoja[hoja])
        if hr is None:
            resultados.append(
                {"seccion": hoja, "estado": "falta", "detalle": f"No se encontró la columna clave '{claves_por_hoja[hoja]}'."}
            )
            continue
        hmap, _ = build_header_map(ws, hr)
        faltantes = [c for c in columnas if c not in hmap]
        if faltantes:
            resultados.append(
                {"seccion": hoja, "estado": "revisar", "detalle": f"No se encontraron columnas: {', '.join(faltantes)}"}
            )
        else:
            resultados.append({"seccion": hoja, "estado": "ok", "detalle": "Todas las columnas esperadas están presentes."})

    return resultados


def mostrar_resultado_validacion(resultados):
    with st.sidebar.expander("📝 Resultado de la validación del archivo", expanded=True):
        for r in resultados:
            texto = f"{r['seccion']}: {r['detalle']}"
            if r["estado"] == "ok":
                st.success(texto)
            elif r["estado"] == "revisar":
                st.warning(texto)
            else:
                st.error(texto)


# ---------------------------------------------------------------------------
# Exportar consulta a Excel
# ---------------------------------------------------------------------------


def generar_reporte_excel(resumen: dict, df_dmas=None, df_punt=None, df_acum=None) -> io.BytesIO:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(list(resumen.items()), columns=["Campo", "Valor"]).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        if df_dmas is not None and not df_dmas.empty:
            df_dmas.to_excel(writer, sheet_name="DMAS", index=False)
        if df_punt is not None and not df_punt.empty:
            df_punt.to_excel(writer, sheet_name="PUNTUALIDAD", index=False)
        if df_acum is not None and not df_acum.empty:
            df_acum[["PERIODO", "AGRUPADOR", "CONCEPTO", "IMPORTE"]].rename(
                columns={"PERIODO": "Semana", "AGRUPADOR": "Agrupador", "CONCEPTO": "Concepto", "IMPORTE": "Importe"}
            ).to_excel(writer, sheet_name="ACUMULADO", index=False)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Respaldo completo
# ---------------------------------------------------------------------------


def generar_respaldo_zip() -> io.BytesIO:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for carpeta in (DATA_DIR, HIST_DIR):
            for archivo in carpeta.rglob("*"):
                if archivo.is_file():
                    zf.write(archivo, archivo.relative_to(BASE_DIR))
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Carga del libro activo (cacheada; se invalida si cambia el archivo)
# ---------------------------------------------------------------------------


@st.cache_resource(show_spinner="Leyendo archivo Excel...")
def load_workbook_data(path_str: str, mtime: float):
    wb = openpyxl.load_workbook(path_str, data_only=True)
    data = {}

    def load_single(sheet_name, key_header_norm):
        ws = wb[sheet_name]
        hr = find_header_row(ws, key_header_norm)
        if hr is None:
            return None
        hmap, cols = build_header_map(ws, hr)
        nom_col = hmap.get(key_header_norm)
        idx = build_nomina_index(ws, hr, nom_col, multi=False)
        return {"ws": ws, "header_row": hr, "hmap": hmap, "columns": cols, "idx": idx}

    def load_multi(sheet_name, key_header_norm):
        ws = wb[sheet_name]
        hr = find_header_row(ws, key_header_norm)
        if hr is None:
            return None
        hmap, cols = build_header_map(ws, hr)
        nom_col = hmap.get(key_header_norm)
        idx = build_nomina_index(ws, hr, nom_col, multi=True)
        return {"ws": ws, "header_row": hr, "hmap": hmap, "columns": cols, "idx": idx}

    data[SHEET_BONOS] = load_single(SHEET_BONOS, "CLAVE") if SHEET_BONOS in wb.sheetnames else None
    data[SHEET_RENDIMIENTO] = (
        load_single(SHEET_RENDIMIENTO, "CLAVE") if SHEET_RENDIMIENTO in wb.sheetnames else None
    )
    data[SHEET_DMAS] = load_single(SHEET_DMAS, "NOMINA") if SHEET_DMAS in wb.sheetnames else None
    data[SHEET_PUNTUALIDAD] = (
        load_multi(SHEET_PUNTUALIDAD, "NO. DE NOMINA") if SHEET_PUNTUALIDAD in wb.sheetnames else None
    )
    return data


# Campos que replican exactamente la pestaña CONSULTA -----------------------
CAMPOS_IZQUIERDA = [
    ("SINIESTROS", SHEET_BONOS, "SINIESTROS"),
    ("QUEJAS", SHEET_BONOS, "QUEJAS"),
    ("SCORE OBTENIDO", SHEET_BONOS, "SCORE OBTENIDO"),
    ("RENDIMIENTO META", SHEET_RENDIMIENTO, "META"),
    ("RENDIMIENTO OBTENIDO", SHEET_RENDIMIENTO, "REDIMIENTO OBTENIDO"),
]
CAMPOS_DERECHA = [
    ("BONO DE SCORE", SHEET_BONOS, "BONO SCORE"),
    ("PUNTUALIDAD", SHEET_BONOS, "PUNTUALIDAD"),
    ("BONO LOGUEO", SHEET_BONOS, "BONO LOGUEO"),
    ("BONO 0 SINIESTROS", SHEET_BONOS, "BONO SINIESTROS"),
    ("BONO 0 QUEJAS", SHEET_BONOS, "BONO QUEJAS"),
    ("BONO SCORE ", SHEET_BONOS, "BONO SCORE1"),
    ("BONO LOGUEO ", SHEET_BONOS, "BONO LOGUEO1"),
    ("BONO RENDIMIENTO", SHEET_RENDIMIENTO, "TOTAL A PAGAR"),
    ("TOTAL BONO DESEMPEÑO", SHEET_BONOS, "TOTAL B.DESEMPEÑO"),
]

COLUMNAS_PUNTUALIDAD = [
    "START_DATE", "ID RUTA", "DES", "START_TIME", "START_ETA", "START_STATUS",
    "END_TIME", "END_ETA", "END_STATUS", "ESTATUS SALIDA", "ESTATUS LLEGADA",
    "ESTATUS", "SEMANA", "PUNTUALIDAD", "TURNO",
]

# ---------------------------------------------------------------------------
# Barra lateral: carga y archivado del archivo semanal
# ---------------------------------------------------------------------------

st.sidebar.header("📂 Archivo semanal")
st.sidebar.caption(
    "Sube el Excel completo de la semana (con todas sus pestañas: BONOS, "
    "RENDIMIENTO, DMAS, PUNTUALIDAD, CONSULTA y ACUMULADO)."
)

meta = json.loads(META_PATH.read_text(encoding="utf-8")) if META_PATH.exists() else None
if meta:
    st.sidebar.success(f"Activo: **{meta['original_filename']}**")
    st.sidebar.caption(f"Cargado el {meta['loaded_at']}")
else:
    st.sidebar.warning("Todavía no se ha cargado ningún archivo.")

if "validacion" in st.session_state:
    mostrar_resultado_validacion(st.session_state["validacion"])

nuevo_archivo = st.sidebar.file_uploader("Subir Excel de la semana (.xlsx)", type=["xlsx"])

if nuevo_archivo is not None:
    if st.sidebar.button("✅ Confirmar carga (archiva la semana anterior)", type="primary"):
        if ACTUAL_PATH.exists() and meta is not None:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            nombre_base = Path(meta["original_filename"]).stem
            archivo_destino = HIST_DIR / f"{nombre_base}_archivado_{stamp}.xlsx"
            shutil.copy2(ACTUAL_PATH, archivo_destino)

        with open(ACTUAL_PATH, "wb") as f:
            f.write(nuevo_archivo.getbuffer())

        META_PATH.write_text(
            json.dumps(
                {
                    "original_filename": nuevo_archivo.name,
                    "loaded_at": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        load_workbook_data.clear()

        # Validar contenido del archivo recién cargado
        st.session_state["validacion"] = validar_archivo(ACTUAL_PATH)

        # El mismo archivo trae la pestaña ACUMULADO: se guarda de una vez.
        grupos = extraer_acumulado_por_periodo(ACTUAL_PATH)
        if grupos:
            guardar_acumulado(grupos)
            st.sidebar.success(
                f"Archivo actualizado. Semana anterior archivada, y se guardó "
                f"ACUMULADO de: {', '.join(grupos.keys())}."
            )
        else:
            st.sidebar.success(
                "Archivo actualizado. La semana anterior quedó en /historico. "
                "(Este archivo no traía pestaña ACUMULADO.)"
            )
        st.rerun()

archivos_historicos = sorted(HIST_DIR.glob("*.xlsx"), reverse=True)
if archivos_historicos:
    with st.sidebar.expander(f"🗂️ Histórico ({len(archivos_historicos)} archivo(s))"):
        for a in archivos_historicos:
            st.write(a.name)

periodos_guardados = listar_periodos_acumulado()
if periodos_guardados:
    with st.sidebar.expander(f"💰 Semanas guardadas en ACUMULADO ({len(periodos_guardados)})"):
        for p in periodos_guardados:
            st.write(p.stem)

# ---------------------------------------------------------------------------
# Respaldo completo
# ---------------------------------------------------------------------------

st.sidebar.markdown("---")
st.sidebar.header("💾 Respaldo")
st.sidebar.caption("Guarda una copia de todo lo cargado (archivo activo, histórico y acumulado).")
if st.sidebar.button("Generar respaldo completo (.zip)"):
    st.session_state["respaldo_buffer"] = generar_respaldo_zip().getvalue()
if "respaldo_buffer" in st.session_state:
    st.sidebar.download_button(
        "⬇️ Descargar respaldo",
        data=st.session_state["respaldo_buffer"],
        file_name=f"respaldo_ConsultaBonosLIPU_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip",
        mime="application/zip",
    )

# ---------------------------------------------------------------------------
# Administrar / eliminar semanas guardadas
# ---------------------------------------------------------------------------

st.sidebar.markdown("---")
with st.sidebar.expander("🗑️ Eliminar una semana completa"):
    st.caption("Esta acción no se puede deshacer.")

    if periodos_guardados:
        st.markdown("**Semana del ACUMULADO (nómina oficial)**")
        opciones_acum = [p.stem for p in periodos_guardados]
        elegido_acum = st.selectbox("Semana a eliminar", opciones_acum, key="sel_acum")
        confirmar_acum = st.checkbox(f"Confirmo eliminar '{elegido_acum}' de ACUMULADO", key="chk_acum")
        if st.button("Eliminar semana de ACUMULADO", disabled=not confirmar_acum):
            (ACUM_DIR / f"{elegido_acum}.csv").unlink(missing_ok=True)
            st.success(f"Semana '{elegido_acum}' eliminada de ACUMULADO.")
            st.rerun()
    else:
        st.caption("No hay semanas de ACUMULADO guardadas todavía.")

    if archivos_historicos:
        st.markdown("**Archivo semanal archivado (BONOS/DMAS/PUNTUALIDAD)**")
        opciones_hist = [a.name for a in archivos_historicos]
        elegido_hist = st.selectbox("Archivo a eliminar", opciones_hist, key="sel_hist")
        confirmar_hist = st.checkbox(f"Confirmo eliminar '{elegido_hist}'", key="chk_hist")
        if st.button("Eliminar archivo histórico", disabled=not confirmar_hist):
            (HIST_DIR / elegido_hist).unlink(missing_ok=True)
            st.success(f"Archivo '{elegido_hist}' eliminado del histórico.")
            st.rerun()
    else:
        st.caption("No hay archivos históricos guardados todavía.")

st.sidebar.markdown("---")
st.sidebar.caption(
    "🔒 Esta carpeta contiene información salarial real. Evita compartirla o "
    "dejarla en una computadora de uso común."
)

# ---------------------------------------------------------------------------
# Cuerpo principal
# ---------------------------------------------------------------------------

st.title("🔎 Consulta de Salarios, Bonos y Puntualidad — LIPU Guadalajara")

if not ACTUAL_PATH.exists():
    st.info("Carga el Excel semanal desde el panel izquierdo para empezar a consultar.")
    st.stop()

data = load_workbook_data(str(ACTUAL_PATH), ACTUAL_PATH.stat().st_mtime)
bonos = data.get(SHEET_BONOS)
rendimiento = data.get(SHEET_RENDIMIENTO)
dmas = data.get(SHEET_DMAS)
puntualidad = data.get(SHEET_PUNTUALIDAD)

tab_individual, tab_multiple, tab_tendencia = st.tabs(
    ["🔎 Consulta individual", "📋 Consulta múltiple", "📈 Tendencia"]
)

# ---------------------------------------------------------------------------
# TAB: Consulta individual (por nómina o por nombre)
# ---------------------------------------------------------------------------

with tab_individual:
    consulta_texto = st.text_input(
        "Número de nómina o nombre del trabajador",
        placeholder="Ej. 14100727 o 'García'",
    )

    nomina = None
    if consulta_texto:
        texto = consulta_texto.strip()
        if texto.isdigit():
            nomina = int(texto)
        else:
            indice_nombres = construir_indice_nombres(bonos, dmas)
            coincidencias = buscar_por_nombre(indice_nombres, texto)
            if not coincidencias:
                st.error(f"No se encontró ningún trabajador cuyo nombre contenga '{texto}'.")
            elif len(coincidencias) == 1:
                nomina, nombre_encontrado = coincidencias[0]
                st.caption(f"Coincidencia única: {nombre_encontrado} (Nómina {nomina})")
            else:
                opciones = [f"{nombre} — Nómina {nom}" for nom, nombre in coincidencias]
                elegido = st.selectbox(
                    f"Se encontraron {len(coincidencias)} coincidencias, elige una:", opciones
                )
                nomina = coincidencias[opciones.index(elegido)][0]

    if nomina is not None:
        fila_bonos = bonos["idx"].get(nomina) if bonos else None
        fila_rendimiento = rendimiento["idx"].get(nomina) if rendimiento else None

        if fila_bonos is None and fila_rendimiento is None:
            # No está en BONOS ni en RENDIMIENTO: se busca como respaldo en ACUMULADO.
            df_acum_respaldo = buscar_acumulado_por_nomina(nomina)
            if df_acum_respaldo.empty:
                st.error("No se encontró esa nómina en BONOS, RENDIMIENTO ni ACUMULADO.")
            else:
                nombre_acum = df_acum_respaldo.iloc[0]["NOMBRE_COMPLETO"]
                st.subheader(f"👤 {nombre_acum} — Nómina {nomina}")
                st.warning(
                    "Esta nómina no aparece en BONOS ni en RENDIMIENTO de la semana "
                    "activa. Se muestra únicamente lo encontrado en ACUMULADO "
                    "(nómina oficial)."
                )
                st.markdown("### 💰 Nómina oficial (ACUMULADO)")
                mostrar_tabla_acumulado(df_acum_respaldo)

                buffer = generar_reporte_excel(
                    {"Nómina": nomina, "Nombre": nombre_acum, "Nota": "No encontrado en BONOS/RENDIMIENTO"},
                    df_acum=df_acum_respaldo,
                )
                st.download_button(
                    "⬇️ Descargar esta consulta en Excel",
                    data=buffer,
                    file_name=f"Consulta_{nomina}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        else:
            nombre = get_val(bonos, fila_bonos, "NOMBRE", default="(nombre no disponible)")
            st.subheader(f"👤 {nombre} — Nómina {nomina}")

            resumen_export = {"Nómina": nomina, "Nombre": nombre}

            col_izq, col_der = st.columns(2)
            with col_izq:
                for etiqueta, hoja, header in CAMPOS_IZQUIERDA:
                    origen = bonos if hoja == SHEET_BONOS else rendimiento
                    fila = fila_bonos if hoja == SHEET_BONOS else fila_rendimiento
                    valor = get_val(origen, fila, header)
                    resumen_export[etiqueta] = to_display(valor)
                    st.metric(etiqueta, to_display(valor))
            with col_der:
                bono_rendimiento_val = None
                for etiqueta, hoja, header in CAMPOS_DERECHA:
                    origen = bonos if hoja == SHEET_BONOS else rendimiento
                    fila = fila_bonos if hoja == SHEET_BONOS else fila_rendimiento
                    valor = get_val(origen, fila, header)
                    resumen_export[etiqueta.strip()] = to_display(valor)
                    st.metric(etiqueta.strip(), to_display(valor))
                    if header == "TOTAL A PAGAR":
                        bono_rendimiento_val = valor
               # st.metric("TOTAL BONO RENDIMIENTO", to_display(bono_rendimiento_val))
                resumen_export["TOTAL BONO RENDIMIENTO"] = to_display(bono_rendimiento_val)

            # ------------------------------------------------------------------
            # DMAS
            # ------------------------------------------------------------------
            df_dmas_export = None
            fila_dmas = dmas["idx"].get(nomina) if dmas else None
            if fila_dmas is not None:
                st.markdown("---")
                st.markdown("### 📋 DMAS")
                registros = [
                    {"Campo": to_display(orig_label), "Valor": to_display(dmas["ws"].cell(row=fila_dmas, column=col).value)}
                    for col, orig_label in dmas["columns"]
                ]
                df_dmas_export = pd.DataFrame(registros)
                st.table(df_dmas_export)
            else:
                st.caption("Sin coincidencias en DMAS para esta nómina.")

            # ------------------------------------------------------------------
            # PUNTUALIDAD
            # ------------------------------------------------------------------
            df_punt_export = None
            filas_punt = puntualidad["idx"].get(nomina, []) if puntualidad else []
            if filas_punt:
                st.markdown("---")
                st.markdown(f"### 🚌 PUNTUALIDAD ({len(filas_punt)} viaje(s) encontrados)")
                registros = []
                for fila in filas_punt:
                    registros.append(
                        {
                            etiqueta_bonita(col_label): to_display(get_val(puntualidad, fila, col_label))
                            for col_label in COLUMNAS_PUNTUALIDAD
                        }
                    )
                df_punt_export = pd.DataFrame(registros)
                st.dataframe(df_punt_export, use_container_width=True, hide_index=True)
            else:
                st.caption("Sin coincidencias en PUNTUALIDAD para esta nómina.")

            # ------------------------------------------------------------------
            # ACUMULADO (nómina oficial)
            # ------------------------------------------------------------------
            df_acum = buscar_acumulado_por_nomina(nomina)
            if not df_acum.empty:
                st.markdown("---")
                st.markdown("### 💰 Nómina oficial (ACUMULADO)")
                mostrar_tabla_acumulado(df_acum)
            else:
                st.caption("Sin coincidencias en ACUMULADO (nómina oficial) para esta nómina.")

            # ------------------------------------------------------------------
            # Exportar
            # ------------------------------------------------------------------
            st.markdown("---")
            buffer = generar_reporte_excel(resumen_export, df_dmas_export, df_punt_export, df_acum)
            st.download_button(
                "⬇️ Descargar esta consulta en Excel",
                data=buffer,
                file_name=f"Consulta_{nomina}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

# ---------------------------------------------------------------------------
# TAB: Consulta múltiple
# ---------------------------------------------------------------------------

with tab_multiple:
    st.subheader("📋 Consulta de varias nóminas a la vez")
    texto_multi = st.text_area(
        "Pega los números de nómina (uno por línea, o separados por coma/espacio)", height=120
    )
    if st.button("Buscar todas", type="primary"):
        crudos = re.split(r"[,\s]+", texto_multi.strip())
        nominas = sorted({int(c) for c in crudos if c.isdigit()})

        if not nominas:
            st.warning("No se detectaron números de nómina válidos.")
        else:
            filas = []
            for nom in nominas:
                fila_b = bonos["idx"].get(nom) if bonos else None
                fila_r = rendimiento["idx"].get(nom) if rendimiento else None
                df_a = buscar_acumulado_por_nomina(nom)
                total_acum = df_a["IMPORTE"].sum() if not df_a.empty else None

                if fila_b is None and fila_r is None:
                    nombre = df_a.iloc[0]["NOMBRE_COMPLETO"] if not df_a.empty else "(no encontrado)"
                    filas.append(
                        {
                            "Nómina": nom,
                            "Nombre": nombre,
                            "Total Bono Desempeño": "",
                            "Total Bono Rendimiento": "",
                            "Total Acumulado": to_display(total_acum) if total_acum is not None else "",
                            "Encontrado en": "Solo ACUMULADO" if not df_a.empty else "No encontrado",
                        }
                    )
                else:
                    nombre = get_val(bonos, fila_b, "NOMBRE", default="")
                    filas.append(
                        {
                            "Nómina": nom,
                            "Nombre": nombre,
                            "Total Bono Desempeño": to_display(get_val(bonos, fila_b, "TOTAL B.DESEMPEÑO")),
                            "Total Bono Rendimiento": to_display(get_val(rendimiento, fila_r, "TOTAL A PAGAR")),
                            "Total Acumulado": to_display(total_acum) if total_acum is not None else "",
                            "Encontrado en": "BONOS/RENDIMIENTO",
                        }
                    )

            df_resultado = pd.DataFrame(filas)
            st.dataframe(df_resultado, use_container_width=True, hide_index=True)

            buffer_multi = io.BytesIO()
            df_resultado.to_excel(buffer_multi, index=False, engine="openpyxl")
            buffer_multi.seek(0)
            st.download_button(
                "⬇️ Descargar tabla en Excel",
                data=buffer_multi,
                file_name="consulta_multiple.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

# ---------------------------------------------------------------------------
# TAB: Tendencia
# ---------------------------------------------------------------------------

with tab_tendencia:
    st.subheader("📈 Tendencia de nómina oficial por semana")
    st.caption("Muestra cómo ha variado lo pagado (ACUMULADO) de un trabajador, semana a semana.")
    texto_tend = st.text_input("Número de nómina", key="tend_nomina", placeholder="Ej. 14100727")

    if texto_tend:
        if texto_tend.strip().isdigit():
            nomina_tend = int(texto_tend.strip())
            df_t = buscar_acumulado_por_nomina(nomina_tend)
            if df_t.empty:
                st.warning("No hay historial de ACUMULADO guardado para esa nómina.")
            else:
                nombre_tend = df_t.iloc[0]["NOMBRE_COMPLETO"]
                st.markdown(f"**{nombre_tend} — Nómina {nomina_tend}**")

                resumen = df_t.groupby("PERIODO")["IMPORTE"].sum().reset_index()
                resumen = resumen.sort_values(by="PERIODO", key=lambda s: s.map(clave_orden_periodo))

                st.bar_chart(resumen.set_index("PERIODO")["IMPORTE"])

                resumen_fmt = resumen.rename(columns={"PERIODO": "Semana", "IMPORTE": "Total pagado"})
                resumen_fmt["Total pagado"] = resumen_fmt["Total pagado"].apply(to_display)
                st.table(resumen_fmt)

                if len(resumen) >= 2:
                    ultimo = resumen.iloc[-1]["IMPORTE"]
                    penultimo = resumen.iloc[-2]["IMPORTE"]
                    diferencia = ultimo - penultimo
                    signo = "subió" if diferencia > 0 else ("bajó" if diferencia < 0 else "se mantuvo igual")
                    st.caption(
                        f"De la semana {resumen.iloc[-2]['PERIODO']} a la {resumen.iloc[-1]['PERIODO']}, "
                        f"el total {signo} ${abs(diferencia):,.2f}."
                    )
        else:
            st.error("Escribe solo el número de nómina (dígitos).")

st.markdown("---")
st.caption(
    "Este documento es de carácter informativo. Su contenido carece de valor oficial; "
    "para el recibo de nómina oficial, solicitarlo en oficinas de Transportes LIPU."
)
